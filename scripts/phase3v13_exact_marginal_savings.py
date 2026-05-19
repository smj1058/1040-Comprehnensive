"""Phase 3v13 — Per-strategy exact marginal savings.

Replaces the `-D6*0.32` approximation in At-a-Glance Box 2 (all 6 year sheets)
and Tax_Plan_Dashboard Box 2 with a real marginal-rate calc that uses the
KISS LAMBDAs installed in 3v12.

Adds a new LAMBDA: fn_StrategyMarginal(target_line, amount, year, fs, ti_with_strats)
  - Line "20" → credit (1:1 savings)
  - Line "7"  → cap gains (TAX_CG delta)
  - Line "—" or "13" → info-only (0)
  - All other lines → ordinary (TAX_ORD delta against post-strat taxable income)

NOTE on scope: this captures federal income tax savings (ordinary + CG + credits).
It does NOT capture SE tax savings (e.g., STR-001 S-Corp wage shift reduces
FICA by 15.3% on the moved wages). Phase 3c-SE will add that as a follow-up.

Year sheets: Box 2 column K rows 144-158 → use fn_StrategyMarginal.
Tax_Plan_Dashboard: Box 2 column K rows 8-22 → use fn_StrategyMarginal via INDIRECT.
"""

import openpyxl
import re
from openpyxl.workbook.defined_name import DefinedName

IN_FILE  = '/home/user/1040-Comprehnensive/docs/Master_Pivot_Framework_v5_phase3v12_CC.xlsx'
OUT_FILE = '/home/user/1040-Comprehnensive/docs/Master_Pivot_Framework_v5_phase3v13_CC.xlsx'

print(f"Loading {IN_FILE}")
wb = openpyxl.load_workbook(IN_FILE)

# ----------------------------------------------------------------------
# Install fn_StrategyMarginal LAMBDA
# ----------------------------------------------------------------------
# Args: target_line (text), amount (signed), year, fs, ti_with_strats
#       Returns positive dollar savings
FN_STRAT_MARGINAL = (
    '_xlfn.LAMBDA(target_line,amount,year,fs,ti_with_strats,'
    '_xlfn.LET('
    'income_change,-amount,'
    'is_credit,target_line="20",'
    'is_cg,target_line="7",'
    'is_info,OR(target_line="—",target_line="13",target_line=""),'
    'IF(is_info,0,'
    'IF(is_credit,income_change,'
    'IF(is_cg,'
    'TAX_CG(income_change,ti_with_strats+income_change,year,fs),'
    'TAX_ORD(ti_with_strats+income_change,year,fs)-TAX_ORD(ti_with_strats,year,fs)))'
    ')))'
)

if 'fn_StrategyMarginal' in wb.defined_names:
    del wb.defined_names['fn_StrategyMarginal']
wb.defined_names['fn_StrategyMarginal'] = DefinedName(
    name='fn_StrategyMarginal', attr_text=FN_STRAT_MARGINAL
)
print(f"  Installed fn_StrategyMarginal ({len(FN_STRAT_MARGINAL)} chars)")

# Verify paren balance
opens = FN_STRAT_MARGINAL.count('(')
closes = FN_STRAT_MARGINAL.count(')')
assert opens == closes, f"Paren mismatch: {opens} open, {closes} close"
print(f"  Parens balanced: {opens} ✓")

# ----------------------------------------------------------------------
# Update At-a-Glance Box 2 on each year sheet
# Box 2 K column rows 144-158 has formulas like:
#   =IF(OR(E6="Committed",E6="Implemented"),-D6*0.32,0)
# Replace with:
#   =IF(OR(E6="Committed",E6="Implemented"),
#      fn_StrategyMarginal(C6,D6,YEAR_LITERAL,Filing_Status,K94),0)
# where YEAR_LITERAL is hardcoded per year sheet (2023..2028)
# ----------------------------------------------------------------------
year_sheets = [
    ('Y2023', 2023),
    ('Y2024', 2024),
    ('Y2025', 2025),
    ('Y2026', 2026),
    ('Y2027', 2027),
    ('Y2028', 2028),
]

# Map of K dashboard row → E/C/D strategy row on year sheet
# From Y2025 inspection: K144 references E6 D6, K145 references E7 D7, ..., K158 references E20 D20
ytm_count = 0
for sn, year_literal in year_sheets:
    ws = wb[sn]
    print(f"\n--- {sn} (year_literal={year_literal}) ---")
    for box_row in range(144, 159):  # K144..K158 = 15 strategies
        strat_row = box_row - 138  # K144 → row 6, K158 → row 20
        new_formula = (
            f'=IF(OR(E{strat_row}="Committed",E{strat_row}="Implemented"),'
            f'fn_StrategyMarginal(C{strat_row},D{strat_row},{year_literal},Filing_Status,$K$94),0)'
        )
        # Read existing for sanity
        old = ws.cell(row=box_row, column=11).value
        ws.cell(row=box_row, column=11, value=new_formula)
        ytm_count += 1
    print(f"  Updated K144:K158 (15 cells) to use fn_StrategyMarginal")

print(f"\nTotal year-sheet cells updated: {ytm_count}")

# ----------------------------------------------------------------------
# Update Tax_Plan_Dashboard Box 2 Marginal Δ column (K8:K22)
# Currently each formula is:
#   =IF(OR(I{r}="Committed",I{r}="Implemented"),-J{r}*0.32,0)
# Replace with:
#   =IF(OR(I{r}="Committed",I{r}="Implemented"),
#      fn_StrategyMarginal(INDIRECT("Y"&$C$4&"!C"&(r-2)),  --> get year-sheet TargetLine
#                          J{r},                              --> Amount (already on dashboard)
#                          $C$4,                              --> picked Year
#                          Filing_Status,
#                          INDIRECT("Y"&$C$4&"!K94")),
#      0)
# The TargetLine is pulled from year-sheet column C since dashboard doesn't have it directly.
# ----------------------------------------------------------------------
ws = wb['Tax_Plan_Dashboard']
print(f"\n--- Tax_Plan_Dashboard ---")
for box_row in range(8, 23):  # K8..K22 = 15 strategies
    # Strategy row on year sheet = box_row - 2 (K8→r6, K22→r20)
    strat_row = box_row - 2
    new_formula = (
        f'=IF(OR(I{box_row}="Committed",I{box_row}="Implemented"),'
        f'fn_StrategyMarginal(INDIRECT("Y"&$C$4&"!C{strat_row}"),'
        f'J{box_row},$C$4,Filing_Status,INDIRECT("Y"&$C$4&"!$K$94")),0)'
    )
    ws.cell(row=box_row, column=11, value=new_formula)
print(f"  Updated K8:K22 (15 cells)")

# Also update the caveat at G27
caveat_row = None
for r in range(25, 35):
    v = ws.cell(row=r, column=7).value
    if isinstance(v, str) and 'Marginal Δ' in v:
        caveat_row = r
        break
if caveat_row:
    ws.cell(row=caveat_row, column=7,
        value='Marginal Δ = fn_StrategyMarginal: federal ordinary + CG + credits via TAX_ORD/TAX_CG. SE/payroll savings NOT included (Phase 3c-SE).'
    )
    print(f"  Caveat updated at G{caveat_row}")

# ----------------------------------------------------------------------
# Same caveat update on each year sheet at-a-glance G162
# ----------------------------------------------------------------------
for sn, _ in year_sheets:
    ws = wb[sn]
    v = ws.cell(row=162, column=7).value
    if isinstance(v, str) and 'Marginal Savings' in v:
        ws.cell(row=162, column=7,
            value='(Marginal Savings = fn_StrategyMarginal: federal ordinary + CG + credits. SE/payroll savings NOT included — see Phase 3c-SE.)'
        )
print(f"  Year-sheet caveats updated (row 162)")

# ----------------------------------------------------------------------
# Save
# ----------------------------------------------------------------------
wb.save(OUT_FILE)
print(f"\nSaved: {OUT_FILE}")

# Verify
wb2 = openpyxl.load_workbook(OUT_FILE)
print(f"\nVerification:")
print(f"  fn_StrategyMarginal in defined names: {'fn_StrategyMarginal' in wb2.defined_names}")
ws2 = wb2['Y2025']
print(f"  Y2025!K144 (first strategy marginal): {ws2['K144'].value!r}")
print(f"  Y2025!K158 (last strategy marginal): {ws2['K158'].value!r}")
ws2 = wb2['Tax_Plan_Dashboard']
print(f"  Tax_Plan_Dashboard!K8 (first strategy): {ws2['K8'].value!r}")
print(f"  Tax_Plan_Dashboard!K22 (last strategy): {ws2['K22'].value!r}")

"""Phase 3v19 — QBI bug fix + audit logic refinement + final verification.

1. Fix QBI bug: J119/K119 on Y2023-Y2028 — change *0.2 to *Rate_QBI_W2_50pct
   so the formula matches the cell label "50% of W-2 Wages" and the IRC
   §199A(b)(2)(B)(i) limit.

2. Rebuild Audit_Report with improved literal extraction that strips
   named ranges before number-finding. Eliminates the 18 "threshold"
   false positives that were really just substrings of named ranges
   (Sec1250, 9999999 bracket upper bounds, etc.).

3. Re-verify everything: LAMBDAs, structured refs, math.
"""

import openpyxl
import re
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.workbook.defined_name import DefinedName
from collections import Counter

IN  = '/home/user/1040-Comprehnensive/docs/Master_Pivot_Framework_v5_phase3v18_CC.xlsx'
OUT = '/home/user/1040-Comprehnensive/docs/Master_Pivot_Framework_v5_phase3v19_CC.xlsx'

wb = openpyxl.load_workbook(IN)

# ----------------------------------------------------------------------
# 1. Fix QBI bug
# ----------------------------------------------------------------------
year_sheets = ['Y2023', 'Y2024', 'Y2025', 'Y2026', 'Y2027', 'Y2028']
qbi_fixes = 0
for sn in year_sheets:
    ws = wb[sn]
    for col in [10, 11]:  # J, K
        for row in [119]:
            cell = ws.cell(row=row, column=col)
            v = cell.value
            if isinstance(v, str) and '*0.2' in v and 'Rate_' not in v:
                cell.value = v.replace('*0.2', '*Rate_QBI_W2_50pct')
                qbi_fixes += 1
print(f"QBI bug fixed in {qbi_fixes} cells (expected 12)")

# ----------------------------------------------------------------------
# 2. Rebuild Audit_Report with improved logic
# ----------------------------------------------------------------------
EXCLUDE = {
    'Tax_Ref', 'Dropdown_Lists', 'Framework_Ref', 'Glossary', 'Setup_Guide',
    'Master_Strategies', 'Property_Appendix', 'Carryovers', 'AsFiled',
    'Master_Inputs', 'Change_Log', 'Connections', 'Projection_History',
    'Control_Panel', 'Audit_Report', 'LAMBDA_Test',
}
OK_INTS = set(range(-2, 31)) | {100, 1000}
OK_YEARS = set(range(2020, 2031))
OK_FLOATS = {0.5, 0.9235}
OK_ENUMS = {'MFJ', 'Single', 'HOH', 'MFS',
            'Committed', 'Implemented', 'Approved', 'Proposed',
            'Baseline', 'Strategy', 'Active', 'Passive', 'Portfolio',
            'Yes', 'No', 'N/A',
            '1z', '2a', '2b', '3a', '3b', '4b', '5b', '6b', '7', '8',
            '—', '', ' '}

def parse_literals_clean(formula):
    """Better literal extraction — strips named ranges before number-finding."""
    s = re.sub(r'\$?[A-Z]{1,3}\$?\d+', ' ', formula)        # cell refs
    s = re.sub(r'R\[?-?\d+\]?C\[?-?\d+\]?', ' ', s)         # R1C1
    s = re.sub(r'\[[^\]]*\]', ' ', s)                       # table [col] refs
    # Strip named-range identifiers (alphabetic words possibly with underscores/digits inside)
    s = re.sub(r'\b[A-Za-z_][A-Za-z0-9_]*\b', ' ', s)
    # Now extract pure numeric literals (guarded so we don't pick numbers inside removed names)
    nums = re.findall(r'-?\d+\.?\d*', s)
    out = []
    for n in nums:
        try:
            f = float(n)
            out.append(int(f) if f == int(f) else f)
        except ValueError: continue
    return out

def parse_strings(formula):
    return re.findall(r'"([^"]*)"', formula)

findings = []
sheets_scanned, cells_scanned = 0, 0
for sn in wb.sheetnames:
    if sn in EXCLUDE: continue
    sheets_scanned += 1
    ws = wb[sn]
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if v is None: continue
            cells_scanned += 1
            coord = cell.coordinate
            if isinstance(v, str) and v in ('#N/A', '#NAME?', '#REF!', '#VALUE!', '#DIV/0!', '#NULL!', '#NUM!'):
                findings.append((sn, coord, 'FLAG', 'Error value', f'Cell = {v}'))
                continue
            if not (isinstance(v, str) and v.startswith('=')): continue
            if re.search(r"'\[?[^']*\.xlsx", v):
                findings.append((sn, coord, 'FLAG', 'External link', f'Ext: {v[:80]}'))
            nums = parse_literals_clean(v)
            for n in nums:
                if isinstance(n, int) and (n in OK_INTS or n in OK_YEARS): continue
                if isinstance(n, float) and n in OK_FLOATS: continue
                if isinstance(n, float) and 0 < n < 1:
                    findings.append((sn, coord, 'WARN', 'Hardcoded rate',
                        f'Rate {n}. Formula: {v[:100]}'))
                elif isinstance(n, int) and abs(n) >= 1000:
                    findings.append((sn, coord, 'WARN', 'Hardcoded threshold',
                        f'Int {n}. Formula: {v[:100]}'))
                elif isinstance(n, float):
                    findings.append((sn, coord, 'WARN', 'Hardcoded decimal',
                        f'Decimal {n}. Formula: {v[:100]}'))
            conds = re.findall(r'IF\([^,)]*[<>]\s*(\d+\.?\d*)', v)
            for c in conds:
                try:
                    n = float(c); n = int(n) if n == int(n) else n
                    if n in OK_INTS or n in OK_YEARS or n in OK_FLOATS: continue
                    findings.append((sn, coord, 'FLAG', 'Magic conditional',
                        f'IF threshold {n}. Formula: {v[:100]}'))
                except ValueError: pass
            if 'INDIRECT(' not in v:
                for s in parse_strings(v):
                    if s in OK_ENUMS or s == '': continue
                    if re.match(r'^Y\d{4}$', s): continue
                    if len(s) == 2 and s.isupper() and s.isalpha(): continue
                    if re.match(r'^STR-\d+$', s): continue
                    if s in ('☑','▣','☐','─') or s.startswith('!'): continue
                    if s in wb.sheetnames: continue
                    if re.match(r'^[!#$:]', s): continue
                    if 'drill' in s.lower() or '↗' in s or '▾' in s or '▸' in s: continue
                    if s.lower() in ('yyyy-mm-dd','mm/dd/yyyy','0.00','#,##0','0%'): continue
                    if len(s) < 3 and not s.isupper(): continue
                    findings.append((sn, coord, 'WARN', 'String literal in formula',
                        f'String {s!r}. Formula: {v[:80]}'))

sev_count = Counter(f[2] for f in findings)
kind_count = Counter(f[3] for f in findings)
print(f"\nNew audit: {len(findings)} findings ({dict(sev_count)})")
print(f"By kind: {dict(kind_count)}")

# Rebuild Audit_Report sheet
if 'Audit_Report' in wb.sheetnames:
    del wb['Audit_Report']
ws = wb.create_sheet('Audit_Report')
wb.move_sheet('Audit_Report', offset=1 - wb.sheetnames.index('Audit_Report'))
ws.sheet_properties.tabColor = 'C00000'

title_font = Font(name='Calibri', size=18, bold=True, color='C00000')
subtitle_font = Font(name='Calibri', size=10, italic=True, color='7F7F7F')
section_font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
section_fill = PatternFill('solid', fgColor='1F4E79')
hdr_font = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
hdr_fill = PatternFill('solid', fgColor='4472C4')
label_font = Font(name='Calibri', size=10)
warn_font = Font(name='Calibri', size=10, color='BF8F00')
warn_fill = PatternFill('solid', fgColor='FFF2CC')
flag_font = Font(name='Calibri', size=10, bold=True, color='C00000')
flag_fill = PatternFill('solid', fgColor='FCE4D6')
ok_font = Font(name='Calibri', size=12, bold=True, color='548235')
ok_fill = PatternFill('solid', fgColor='E2EFDA')
center = Alignment(horizontal='center', vertical='center')
left = Alignment(horizontal='left', vertical='center', wrap_text=True)
thin = Border(left=Side('thin', color='BFBFBF'), right=Side('thin', color='BFBFBF'),
              top=Side('thin', color='BFBFBF'), bottom=Side('thin', color='BFBFBF'))

ws['B2'] = 'AUDIT REPORT — v19'
ws['B2'].font = title_font
ws['B3'] = (f'Static scan of {sheets_scanned} sheets / {cells_scanned:,} cells. '
            f'v19 improvements: literal extraction now strips named ranges before number-finding '
            f'(eliminates ~18 false-positive "threshold" findings from prior runs).')
ws['B3'].font = subtitle_font
ws.row_dimensions[2].height = 26

# Resolution log
ws.merge_cells('B5:G5')
ws['B5'] = '✅ RESOLVED IN v18-v19'
ws['B5'].font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
ws['B5'].fill = PatternFill('solid', fgColor='548235')
ws['B5'].alignment = left

resolved = [
    'v18: 11 rate names defined (Rate_FICA_*, Rate_QBI_*, Rate_NIIT, Rate_SE_*, Rate_Sec1250_*)',
    'v18: 72 year-sheet cells substituted literals → named references',
    'v18: fn_StrategyMarginal, TAX_NIIT, TAX_SE LAMBDAs updated to use named rates',
    'v19: QBI bug fixed — J119/K119 on all 6 year sheets: *0.2 → *Rate_QBI_W2_50pct',
    'v19: 18 "threshold" findings confirmed as false positives from named-range substrings — audit logic refined',
]
for i, r in enumerate(resolved):
    rr = 7 + i
    ws.cell(row=rr, column=2, value='•').font = Font(name='Calibri', size=10, bold=True, color='548235')
    ws.cell(row=rr, column=2).alignment = center
    ws.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=7)
    ws.cell(row=rr, column=3, value=r).font = label_font
    ws.cell(row=rr, column=3).alignment = left
    ws.row_dimensions[rr].height = 22

# Summary
sum_start = 7 + len(resolved) + 2
ws.merge_cells(start_row=sum_start, start_column=2, end_row=sum_start, end_column=7)
ws.cell(row=sum_start, column=2, value='▸ SECTION A — SUMMARY').font = section_font
ws.cell(row=sum_start, column=2).fill = section_fill
ws.cell(row=sum_start, column=2).alignment = left

stats = [
    ('Sheets scanned', sheets_scanned),
    ('Cells scanned', cells_scanned),
    ('🔴 FLAG (high-confidence bug)', sev_count.get('FLAG', 0)),
    ('🟡 WARN (review)', sev_count.get('WARN', 0)),
]
for i, (lbl, val) in enumerate(stats):
    r = sum_start + 2 + i
    ws.cell(row=r, column=2, value=lbl).font = label_font
    ws.cell(row=r, column=3, value=val).font = Font(name='Calibri', size=10, bold=True)
    ws.cell(row=r, column=3).alignment = center

# Breakdown
bkd_start = sum_start + 2 + len(stats) + 2
ws.merge_cells(start_row=bkd_start, start_column=2, end_row=bkd_start, end_column=7)
ws.cell(row=bkd_start, column=2, value='▸ SECTION B — BREAKDOWN BY KIND').font = section_font
ws.cell(row=bkd_start, column=2).fill = section_fill
ws.cell(row=bkd_start, column=2).alignment = left
for i, (k, c) in enumerate(kind_count.most_common()):
    r = bkd_start + 2 + i
    ws.cell(row=r, column=2, value=k).font = label_font
    ws.cell(row=r, column=3, value=c).font = Font(name='Calibri', size=10, bold=True)
    ws.cell(row=r, column=3).alignment = center

# Findings table (top 500)
find_start = bkd_start + 2 + len(kind_count) + 2
ws.merge_cells(start_row=find_start, start_column=2, end_row=find_start, end_column=7)
ws.cell(row=find_start, column=2,
    value=f'▸ SECTION C — FINDINGS ({len(findings)} total)').font = section_font
ws.cell(row=find_start, column=2).fill = section_fill
ws.cell(row=find_start, column=2).alignment = left

hdr_r = find_start + 2
for i, h in enumerate(['Severity','Sheet','Cell','Kind','Detail','Action']):
    c = ws.cell(row=hdr_r, column=2+i, value=h)
    c.font = hdr_font; c.fill = hdr_fill; c.alignment = center; c.border = thin

sev_order = {'FLAG':0, 'WARN':1}
sorted_f = sorted(findings, key=lambda x: (sev_order.get(x[2], 2), x[0], x[1]))
actions = {
    'Hardcoded rate': 'Define a named rate or pull from tblThresholds.',
    'Hardcoded threshold': 'Replace with named range or tblThresholds lookup.',
    'Hardcoded decimal': 'Verify if literal should be named.',
    'String literal in formula': 'If enum, pull from Dropdown_Lists.',
    'Magic conditional': 'Replace literal in IF with named range / tblThresholds.',
    'Error value': 'Investigate broken dependency.',
    'External link': 'Confirm intentional.',
}
MAX = 500
for i, (sn, coord, sev, kind, detail) in enumerate(sorted_f[:MAX]):
    r = hdr_r + 1 + i
    sc = ws.cell(row=r, column=2, value=sev)
    sc.font = flag_font if sev=='FLAG' else warn_font
    sc.fill = flag_fill if sev=='FLAG' else warn_fill
    sc.alignment = center; sc.border = thin
    for col, val in enumerate([sn, coord, kind, detail, actions.get(kind, '')], start=3):
        c = ws.cell(row=r, column=col, value=val)
        c.font = Font(name='Calibri', size=9) if col>3 else label_font
        c.alignment = left if col>=5 else center
        c.border = thin
if len(sorted_f) > MAX:
    r = hdr_r + 1 + MAX
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
    ws.cell(row=r, column=2, value=f'… {len(sorted_f) - MAX} more findings truncated.').font = subtitle_font

for col, width in zip(['A','B','C','D','E','F','G'], [3, 10, 22, 10, 22, 70, 50]):
    ws.column_dimensions[col].width = width

# ----------------------------------------------------------------------
# 3. Save + verify
# ----------------------------------------------------------------------
wb.save(OUT)
print(f"\nSaved: {OUT}")

# Verify QBI fix took
wb2 = openpyxl.load_workbook(OUT)
ws = wb2['Y2025']
print(f"\nQBI fix verification (Y2025):")
print(f"  J119: {ws['J119'].value}")
print(f"  K119: {ws['K119'].value}")

# LAMBDA test re-check
print(f"\nLAMBDA defined names:")
for name in ['TAX_ORD','TAX_CG','TAX_NIIT','TAX_SE','STD_DED','TAX_STATE','fn_StrategyMarginal']:
    if name in wb2.defined_names:
        v = wb2.defined_names[name].value
        print(f"  ✓ {name} ({len(v)} chars)")
    else:
        print(f"  ✗ {name} MISSING")

# Rate names check
print(f"\nRate names:")
for name in ['Rate_FICA_SS_Employee','Rate_FICA_Combined','Rate_QBI_Deduction','Rate_QBI_W2_50pct','Rate_NIIT']:
    if name in wb2.defined_names:
        print(f"  ✓ {name} = {wb2.defined_names[name].value}")
    else:
        print(f"  ✗ {name} MISSING")
